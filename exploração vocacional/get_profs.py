
import docx # type: ignore
import json
import re
from typing import Any, cast

doc = docx.Document('EXPLORAÇÃO VOCACIONAL explicação.docx')
professions = []
current_group = None

# We know professions are at the end.
# They seem to be grouped.
group_pattern = re.compile(r'^\d+\.\s+(.*)')

for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if not text:
        continue
    
    # Check for group headers
    match_group = group_pattern.match(text)
    if match_group and len(text) < 50:
        current_group = match_group.group(1)
        continue
    
    # Look for profession entries
    # Heuristic: the profession name is usually a single line, 
    # and the next line has "Formação:" or "Descrição:"
    if i < len(doc.paragraphs) - 1:
        next_text = doc.paragraphs[i+1].text.strip()
        if "Formação:" in next_text or "Descrição:" in next_text or "Ambientes:" in next_text:
            name = text
            details = {}
            
            # Extract details from subsequent paragraphs until we hit a new profession or group
            j = i + 1
            while j < len(doc.paragraphs):
                detail_text = doc.paragraphs[j].text.strip()
                if not detail_text:
                    j += 1
                    continue
                
                # If we hit something that looks like a new profession or group, stop
                detail_sample = ""
                for char in cast(Any, detail_text):
                    if len(detail_sample) >= 20: break
                    detail_sample += char
                if (j > i + 1 and not (":" in detail_sample)) or group_pattern.match(detail_text):
                    break
                
                if ":" in detail_text:
                    key, val = detail_text.split(":", 1)
                    details[key.strip()] = val.strip()
                
                j += 1
            
            professions.append({
                "name": name,
                "group": current_group,
                **details
            })

with open('professions_list.json', 'w', encoding='utf-8') as f:
    json.dump(professions, f, indent=2, ensure_ascii=False)

print(f"Extracted {len(professions)} professions.")
for i, p in enumerate(professions):
    if i >= 5: break
    print(f"- {p['name']} ({p.get('group', 'N/A')})")
